"""从简历文件抽出纯文本（不做 OCR）。"""

from __future__ import annotations

import io
import re
from typing import Tuple


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024  # 8MB


def sniff_ext(filename: str) -> str:
    name = (filename or "").lower().strip()
    for ext in ALLOWED_EXTENSIONS:
        if name.endswith(ext):
            return ext
    return ""


def extract_text(filename: str, data: bytes) -> Tuple[str, str]:
    """
    返回 (text, warning)。
    warning 非空时表示抽字可能不完整（如扫描件）。
    """
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError(f"文件过大（上限 {MAX_UPLOAD_BYTES // (1024 * 1024)}MB）")

    ext = sniff_ext(filename)
    if not ext:
        raise ValueError("仅支持 PDF / Word(.docx) / 纯文本(.txt)")

    if ext == ".txt":
        text = data.decode("utf-8", errors="ignore")
    elif ext == ".docx":
        text = _from_docx(data)
    else:
        text = _from_pdf(data)

    text = _normalize(text)
    warning = ""
    if len(text) < 40:
        warning = (
            "几乎抽不到文字。若是扫描件/图片 PDF，请换可选中文字的 PDF，"
            "或直接粘贴简历文本。"
        )
    return text, warning


def _normalize(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError("服务器未安装 python-docx，无法解析 Word") from e

    doc = Document(io.BytesIO(data))
    parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    # 表格里的经历也常见
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ValueError("服务器未安装 pypdf，无法解析 PDF") from e

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        try:
            t = page.extract_text() or ""
        except Exception:
            t = ""
        if t.strip():
            parts.append(t.strip())
    return "\n\n".join(parts)
